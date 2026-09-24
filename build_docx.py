# -*- coding: utf-8 -*-
"""build_docx.py — 将 manuscript/ 下的全部章节合并为一个独立 Word 文档。
用法：python build_docx.py [输出文件名]
依赖：python-docx（安装到 WorkBuddy 托管 venv）
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
MANUSCRIPT = os.path.join(BASE, "manuscript")

CHAPTER_FILES = [
    "00-前言.md",
    "01-范式转移.md",
    "02-认知重构.md",
    "03-提问的艺术.md",
    "04-学习革命.md",
    "05-专业深耕.md",
    "06-创作与产出.md",
    "07-学术规范.md",
    "08-职业竞争力.md",
    "09-生活心理边界.md",
    "10-伦理与未来.md",
    "11-后记.md",
]
APPENDIX_FILES = [
    "appendix/A-工具速查表.md",
    "appendix/B-提示词模板库.md",
    "appendix/C-高校AI政策查询指南.md",
    "appendix/D-延伸阅读.md",
    "appendix/E-术语词典.md",
]

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_east_asia(el, font_name):
    """为 oxml 元素（样式或 run 的底层元素）设置中文字体（eastAsia）。"""
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def init_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(12)
    set_east_asia(st.element, BODY_FONT)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)

    for name, size, bold in [("Heading 1", 22, True), ("Heading 2", 16, True), ("Heading 3", 14, True)]:
        hs = doc.styles[name]
        hs.font.name = HEAD_FONT
        hs.font.size = Pt(size)
        hs.font.bold = bold
        hs.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_east_asia(hs.element, HEAD_FONT)
        hs.paragraph_format.space_before = Pt(18)
        hs.paragraph_format.space_after = Pt(10)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


INLINE_PATTERNS = [
    (re.compile(r"\*\*(.+?)\*\*"), "bold"),
    (re.compile(r"\*(.+?)\*"), "italic"),
    (re.compile(r"`([^`]+)`"), "mono"),
]


def add_runs(par, text):
    """把 markdown 行内标记（**加粗** *斜体* `代码`）转换为 run。"""
    segments = [(text, set())]
    for pat, style in INLINE_PATTERNS:
        out = []
        for seg_text, seg_styles in segments:
            last = 0
            for m in pat.finditer(seg_text):
                out.append((seg_text[last:m.start()], seg_styles))
                out.append((m.group(1), seg_styles | {style}))
                last = m.end()
            out.append((seg_text[last:], seg_styles))
        segments = out
    for seg_text, seg_styles in segments:
        if not seg_text:
            continue
        run = par.add_run(seg_text)
        if "bold" in seg_styles:
            run.bold = True
        if "italic" in seg_styles:
            run.italic = True
        if "mono" in seg_styles:
            run.font.name = MONO_FONT
            run.font.size = Pt(10.5)
            set_east_asia(run._element, BODY_FONT)


def add_table(doc, rows):
    """rows: 二维列表，第一行为表头。"""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.2
            add_runs(par, cell_text)
            for run in par.runs:
                run.font.size = Pt(10.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_toc(doc):
    par = doc.add_paragraph()
    run = par.add_run()
    fld_begin = par._p.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = par._p.makeelement(qn("w:instrText"), {})
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = par._p.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "separate"})
    hint = par._p.makeelement(qn("w:t"), {})
    hint.text = "【在 Word 中打开后：引用 → 目录 → 更新目录，或右键此处选择更新域】"
    fld_end = par._p.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    for el in (fld_begin, instr, fld_sep, hint, fld_end):
        run._r.append(el)


def render_markdown(doc, path):
    with open(path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = []
            for ln in table_buf:
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue  # 分隔行
                rows.append(cells)
            if rows:
                add_table(doc, rows)
            table_buf = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            table_buf.append(line)
            continue
        flush_table()
        if not line.strip():
            continue
        if re.fullmatch(r"-{3,}", line.strip()):
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("> "):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(24)
            par.paragraph_format.first_line_indent = Pt(0)
            add_runs(par, line[2:].strip())
        elif re.fullmatch(r"\*[^*].*\*", line.strip()):
            # 整行斜体（初稿注释等）
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Pt(0)
            run = par.add_run(line.strip().strip("*"))
            run.italic = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        else:
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Pt(24)
            add_runs(par, line.strip())
    flush_table()


def main():
    out_name = sys.argv[1] if len(sys.argv) > 1 else "与AI共生-全书初稿.docx"
    out_path = os.path.join(BASE, out_name)

    doc = Document()
    init_styles(doc)

    # ---- 封面 ----
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("与AI共生")
    r.font.size = Pt(36)
    r.bold = True
    r.font.name = HEAD_FONT
    set_east_asia(r._element, HEAD_FONT)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("AI时代大学生的生成指南")
    r.font.size = Pt(18)
    r.font.name = HEAD_FONT
    set_east_asia(r._element, HEAD_FONT)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(48)
    r = note.add_run("全书初稿 · 各章按简编篇幅合成，扩写至目标字数的工作仍在进行\n文中 [待核实] 标记处需在出版前完成人工核验")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_page_break()

    # ---- 目录 ----
    doc.add_heading("目 录", level=1)
    add_toc(doc)
    doc.add_page_break()

    # ---- 正文 ----
    for rel in CHAPTER_FILES + APPENDIX_FILES:
        path = os.path.join(MANUSCRIPT, rel)
        if not os.path.exists(path):
            print(f"[跳过] 缺失: {rel}")
            continue
        if rel != CHAPTER_FILES[0]:
            doc.add_page_break()
        render_markdown(doc, path)

    doc.save(out_path)
    print(f"[完成] {out_path}")

    # 统计字数
    total = 0
    for rel in CHAPTER_FILES + APPENDIX_FILES:
        path = os.path.join(MANUSCRIPT, rel)
        if os.path.exists(path):
            text = open(path, "r", encoding="utf-8").read()
            cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
            total += cjk
            print(f"  {rel}: {cjk} 汉字")
    print(f"  合计: {total} 汉字")


if __name__ == "__main__":
    main()
